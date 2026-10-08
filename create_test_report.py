"""Create a self-contained formatting fixture; no missing external image needed."""
import base64
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def main():
    output = Path(__file__).parent / 'test-assets' / 'messy-report.docx'
    output.parent.mkdir(exist_ok=True)
    doc = Document()
    doc.add_paragraph('Formatting test report', style='Title')
    heading = doc.add_heading('First lesson', 1); heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i in range(3):
        p = doc.add_paragraph('This deliberately uneven paragraph tests fonts, spacing, alignment and indentation. ' * 3)
        p.alignment = [WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY][i]
        p.paragraph_format.left_indent = Inches(.2 * i)
        p.paragraph_format.line_spacing = .8 + i * .5
        p.runs[0].font.name = ['Arial', 'Courier New', 'Calibri'][i]
        p.runs[0].font.size = Pt([9, 17, 11][i])
    doc.add_paragraph('Numbered item with real Word numbering', style='List Number')
    doc.add_paragraph('Second numbered item', style='List Number')
    table = doc.add_table(rows=3, cols=2); table.style = 'Table Grid'
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = ['Topic', 'Notes'][j] if i == 0 else f'Row {i} column {j + 1}'
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            cell.paragraphs[0].runs[0].font.name = 'Courier New'
    pixel = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL8xQAAAABJRU5ErkJggg==')
    doc.add_picture(BytesIO(pixel), width=Inches(7.2), height=Inches(.25))
    doc.add_paragraph('Figure 1. Oversized test image', style='Caption')
    doc.sections[0].header.paragraphs[0].text = 'Existing header'
    doc.sections[0].footer.paragraphs[0].text = 'Existing footer'
    doc.save(output)
    print(output)


if __name__ == '__main__': main()
