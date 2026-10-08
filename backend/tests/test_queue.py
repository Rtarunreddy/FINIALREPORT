import os
import shutil
import threading
import time
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi.testclient import TestClient

from app import pdf_export
from app.config import Settings
from app.db_models import Document as Doc, Job, utcnow
from app.formatting import DEFAULT_PROFILE, apply_profile
from app.main import create_app
from app.summary import summarize
from app.worker import JobRunner
from conftest import build, run_apply, sign_up

HAS_LO = bool(shutil.which("soffice") or shutil.which("libreoffice"))


def docx_bytes(text="Queue test."):
    d = Document(); d.add_paragraph(text); s = BytesIO(); d.save(s); return s.getvalue()


def upload(c, text="Queue test."):
    r = c.post("/api/uploads", files={"file": ("r.docx", docx_bytes(text))}); assert r.status_code == 200, r.text; return r.json()["id"]


def queue(c, doc, **body):
    r = c.post(f"/api/uploads/{doc}/apply", json=body); assert r.status_code == 202, r.text; return r.json()["id"]


# ---- queue mechanics ----
def test_apply_returns_immediately_and_runs_later(cloud):
    c = sign_up(cloud.app); doc = upload(c)
    job = queue(c, doc)
    assert c.get(f"/api/jobs/{job}").json()["status"] == "queued"
    assert c.get(f"/api/jobs/{job}").json()["outputId"] is None
    assert cloud.app.state.runner.run_pending() == 1
    done = c.get(f"/api/jobs/{job}").json()
    assert done["status"] == "done" and done["outputId"] and done["report"]["before"] == done["report"]["after"]


def test_each_job_is_claimed_exactly_once_by_competing_workers(cloud):
    c = sign_up(cloud.app); docs = [upload(c) for _ in range(5)]
    cloud.app.state.settings.max_active_jobs = 50
    ids = [queue(c, d) for d in docs]
    runners = [JobRunner(cloud.app.state.settings, cloud.app.state.session_factory, cloud.app.state.storage) for _ in range(3)]
    counts = []
    threads = [threading.Thread(target=lambda r=r: counts.append(r.run_pending())) for r in runners]
    [t.start() for t in threads]; [t.join() for t in threads]
    assert sum(counts) == 5
    with cloud.app.state.session_factory() as db:
        jobs = [db.get(Job, i) for i in ids]
        assert all(j.status == "done" and j.attempts == 1 for j in jobs)
        assert len({j.output_doc_id for j in jobs}) == 5


def test_real_worker_threads_finish_jobs(tmp_path):
    with TestClient(build(tmp_path, workers=2)) as base:
        c = sign_up(base.app); job = queue(c, upload(c))
        for _ in range(100):
            status = c.get(f"/api/jobs/{job}").json()
            if status["status"] in {"done", "failed"}: break
            time.sleep(0.1)
        assert status["status"] == "done", status


def test_failed_job_reports_a_clear_error_and_leaves_no_output(cloud):
    c = sign_up(cloud.app); doc = upload(c)
    job = queue(c, doc)
    with cloud.app.state.session_factory() as db:
        source = db.get(Doc, doc); key = source.storage_key
    cloud.app.state.storage.delete(key)  # file vanished from storage
    cloud.app.state.runner.run_pending()
    view = c.get(f"/api/jobs/{job}").json()
    assert view["status"] == "failed" and view["error"] and view["outputId"] is None
    with cloud.app.state.session_factory() as db: assert db.query(Doc).filter(Doc.kind == "output").count() == 0


def test_job_for_deleted_or_expired_upload_fails_cleanly(cloud):
    c = sign_up(cloud.app); doc = upload(c); job = queue(c, doc)
    with cloud.app.state.session_factory() as db:
        row = db.get(Doc, doc); row.expires_at = utcnow().replace(year=2000); db.commit()
    cloud.app.state.runner.run_pending()
    view = c.get(f"/api/jobs/{job}").json()
    assert view["status"] == "failed" and "no longer available" in view["error"]


def test_deleting_an_upload_cancels_its_queued_jobs(cloud):
    c = sign_up(cloud.app); doc = upload(c); job = queue(c, doc)
    assert c.delete(f"/api/uploads/{doc}").status_code == 200
    assert cloud.app.state.runner.run_pending() == 0
    view = c.get(f"/api/jobs/{job}").json(); assert view["status"] == "failed" and "Cancelled" in view["error"]


def test_stale_running_jobs_are_recovered(cloud):
    c = sign_up(cloud.app); job = queue(c, upload(c))
    with cloud.app.state.session_factory() as db:
        row = db.get(Job, job); row.status, row.started_at = "running", utcnow().replace(year=2020); db.commit()
    assert cloud.app.state.runner.recover_stale() == 1
    assert c.get(f"/api/jobs/{job}").json()["status"] == "failed"


def test_active_job_limit_per_user(cloud):
    c = sign_up(cloud.app); doc = upload(c)
    cloud.app.state.settings.max_active_jobs = 2
    queue(c, doc); queue(c, doc)
    assert c.post(f"/api/uploads/{doc}/apply", json={}).status_code == 429
    other = sign_up(cloud.app, "other@example.com")
    assert other.post(f"/api/uploads/{upload(other)}/apply", json={}).status_code == 202   # limit is per user
    cloud.app.state.runner.run_pending()
    assert c.post(f"/api/uploads/{doc}/apply", json={}).status_code == 202


def test_job_history_is_private_and_light(cloud):
    alice, bob = sign_up(cloud.app, "alice@example.com"), sign_up(cloud.app, "bob@example.com")
    run_apply(alice, upload(alice))
    rows = alice.get("/api/jobs").json()
    assert len(rows) == 1 and rows[0]["status"] == "done" and rows[0]["report"] is None
    assert bob.get("/api/jobs").json() == []
    assert bob.get(f"/api/jobs/{rows[0]['id']}").status_code == 404
    assert bob.get("/api/jobs/not-a-uuid").status_code == 404


# ---- one-click default and templates ----
def test_one_click_uses_default_template_without_a_body(cloud):
    c = sign_up(cloud.app); result = run_apply(c, upload(c))
    assert result["status"] == "done" and "Times New Roman" in result["report"]["formatting_preview"]["body"]


def test_template_id_selects_builtin_or_own_template_only(cloud):
    alice, bob = sign_up(cloud.app, "alice@example.com"), sign_up(cloud.app, "bob@example.com")
    ieee = run_apply(alice, upload(alice), templateId="ieee")
    assert ieee["status"] == "done" and "10 pt" in ieee["report"]["formatting_preview"]["body"]
    mine = alice.post("/api/profiles", json={**DEFAULT_PROFILE, "id": "", "name": "Mine"}).json()["id"]
    assert run_apply(alice, upload(alice), templateId=mine)["status"] == "done"
    with cloud.app.state.session_factory() as db: assert db.query(Job).filter(Job.template_id == mine).count() == 1
    assert bob.post(f"/api/uploads/{upload(bob)}/apply", json={"templateId": mine}).status_code == 404
    assert alice.post(f"/api/uploads/{upload(alice)}/apply", json={"templateId": "missing"}).status_code == 404


# ---- before/after summary ----
def messy_document(path):
    d = Document()
    for i, (font, size, align) in enumerate([("Arial", 10, WD_ALIGN_PARAGRAPH.LEFT), ("Calibri", 12, WD_ALIGN_PARAGRAPH.JUSTIFY),
                                             ("Verdana", 9, WD_ALIGN_PARAGRAPH.CENTER), ("Arial", 14, WD_ALIGN_PARAGRAPH.RIGHT)]):
        p = d.add_paragraph(); r = p.add_run(f"Paragraph {i} with some words in it to give weight."); r.font.name, r.font.size = font, Pt(size); p.alignment = align
    d.add_heading("Chapter", 1).runs[0].font.name = "Comic Sans MS"
    d.save(path)


def test_summary_reports_unification(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"; messy_document(src)
    apply_profile(src, out, DEFAULT_PROFILE)
    s = summarize(src, out); rows = {r["label"]: r for r in s["rows"]}
    assert rows["Body fonts"]["before"].count(",") >= 2 and rows["Body fonts"]["after"] == "Times New Roman"
    assert rows["Body sizes"]["after"] == "12 pt" and rows["Body alignment"]["after"] == "justified"
    assert any(h.startswith("Different body fonts: ") and h.endswith("→ 1") for h in s["highlights"])
    assert s["bodyCombinations"]["before"] == 4 and s["bodyCombinations"]["after"] == 1
    assert rows["Heading 1"]["before"].startswith("Comic Sans MS") and rows["Heading 1"]["after"].startswith("Times New Roman")


def test_summary_handles_documents_without_body_text(tmp_path):
    src = tmp_path / "e.docx"; Document().save(src)
    s = summarize(src, src); assert s["bodyCombinations"] == {"before": 0, "after": 0} and s["highlights"] == []


def test_job_report_includes_summary(cloud):
    c = sign_up(cloud.app); result = run_apply(c, upload(c))
    assert "rows" in result["report"]["summary"] and any(r["label"] == "Page margins" for r in result["report"]["summary"]["rows"])


# ---- PDF export ----
def test_pdf_unavailable_gives_clear_error(tmp_path):
    with TestClient(build(tmp_path, soffice_path="/nonexistent/soffice")) as base:
        c = sign_up(base.app); out = run_apply(c, upload(c))["id"]
        if pdf_export.word_available(): pytest.skip("Word is available")
        r = c.post(f"/api/outputs/{out}/pdf"); assert r.status_code == 400 and "not available" in r.json()["detail"]
        assert c.get("/api/health").json()["pdfExport"] is False


@pytest.mark.skipif(not HAS_LO, reason="LibreOffice not installed")
def test_pdf_export_end_to_end_with_reuse_and_cleanup(cloud):
    c = sign_up(cloud.app); doc = upload(c); out = run_apply(c, doc)["id"]
    assert c.get("/api/health").json()["pdfExport"] is True
    r = c.post(f"/api/outputs/{out}/pdf"); assert r.status_code == 202
    again = c.post(f"/api/outputs/{out}/pdf"); assert again.json()["id"] == r.json()["id"]      # no duplicate while queued
    cloud.app.state.runner.run_pending()
    job = c.get(f"/api/jobs/{r.json()['id']}").json()
    assert job["status"] == "done" and job["report"]["engine"] == "libreoffice" and job["report"]["pages"] >= 1
    pdf = c.get(f"/api/outputs/{job['outputId']}/document")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-" and pdf.headers["content-type"] == "application/pdf"
    assert "formatted-report.pdf" in pdf.headers["content-disposition"]
    assert c.post(f"/api/outputs/{out}/pdf").json()["id"] == r.json()["id"]                       # finished export is reused
    assert c.post(f"/api/outputs/{job['outputId']}/pdf").status_code == 400                       # PDFs are not re-exported
    assert c.delete(f"/api/uploads/{doc}").status_code == 200
    assert c.get(f"/api/outputs/{job['outputId']}/document").status_code == 404                   # PDF removed with its upload


def fake_soffice(tmp_path, body):
    path = tmp_path / "soffice"; path.write_text("#!/bin/sh\n" + body); path.chmod(0o755); return str(path)


@pytest.mark.skipif(os.name == "nt", reason="shell script stand-ins")
def test_hung_or_garbage_converter_fails_cleanly(tmp_path):
    src = tmp_path / "a.docx"; src.write_bytes(docx_bytes())
    start = time.time()
    with pytest.raises(ValueError, match="too long"): pdf_export.libreoffice_convert(fake_soffice(tmp_path, "sleep 30\n"), src, tmp_path, 1)
    assert time.time() - start < 10
    (tmp_path / "g").mkdir()
    garbage = fake_soffice(tmp_path / "g", 'while [ $# -gt 0 ]; do [ "$1" = "--outdir" ] && out="$2"; shift; done\necho notapdf > "$out/a.pdf"\n')
    with pytest.raises(ValueError, match="could not be created"): pdf_export.libreoffice_convert(garbage, src, tmp_path, 5)
    assert (tmp_path / "pdf-out" / "a.pdf").read_bytes() == b"notapdf\n"   # the stand-in did write a file; it was rejected for its content
