import base64
import json
from copy import deepcopy
from io import BytesIO
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.opc.constants import RELATIONSHIP_TYPE
from fastapi.testclient import TestClient
import pytest

from app.db_models import Feedback
from conftest import run_apply, stored_files
from app.formatting import DEFAULT_PROFILE, apply_profile, signature, metrics, style_profile, column_width


def document_bytes():
    doc = Document(); doc.add_paragraph("A short report."); stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


def upload(client, payload=None):
    response = client.post("/api/uploads", files={"file": ("report.docx", payload or document_bytes())})
    assert response.status_code == 200, response.text
    return response.json()["id"]


def test_api_complete_flow_and_deletion(client):
    ident = upload(client)
    assert client.get("/api/health").json()["status"] == "ok"
    assert client.post(f"/api/uploads/{ident}/audit", json={"profile": DEFAULT_PROFILE}).status_code == 200
    result = run_apply(client, ident, DEFAULT_PROFILE)
    assert result["status"] == "done", result
    assert result["report"]["before"] == result["report"]["after"]
    assert Document(BytesIO(client.get(f"/api/outputs/{result['id']}/document").content)).paragraphs[0].text == "A short report."
    assert client.get(f"/api/outputs/{result['id']}/report").status_code == 200
    assert client.delete(f"/api/uploads/{ident}").status_code == 200
    assert client.get(f"/api/outputs/{result['id']}/document").status_code == 404
    assert client.post(f"/api/uploads/{ident}/apply", json={"profile": DEFAULT_PROFILE}).status_code == 404


def test_review_is_validated_and_saved(client):
    assert client.post("/api/reviews", json={"rating": 0}).status_code == 422
    assert client.post("/api/reviews", json={"rating": 5, "comment": "Very useful"}).status_code == 200
    with client.app.state.session_factory() as db: assert db.query(Feedback).one().comment == "Very useful"


@pytest.mark.parametrize("payload", [b"not a document", b"PKbad zip"])
def test_bad_upload_does_not_leave_files(client, payload):
    response = client.post("/api/uploads", files={"file": ("bad.docx", payload)})
    assert response.status_code == 400
    assert not stored_files(client)


def test_zip_traversal_is_rejected(client):
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as z:
        z.writestr("..\\outside", "x"); z.writestr("word/document.xml", "x")
    assert client.post("/api/uploads", files={"file": ("bad.docx", stream.getvalue())}).status_code == 400
    assert not stored_files(client)


def test_oversized_upload_is_removed(client):
    client.app.state.settings.max_upload = 12
    assert client.post("/api/uploads", files={"file": ("big.docx", document_bytes())}).status_code == 400
    assert not stored_files(client)


@pytest.mark.parametrize("field,value", [("size", -1), ("alignment", "diagonal"), ("lineSpacing", 99)])
def test_invalid_profile_returns_actionable_validation(client, field, value):
    profile = deepcopy(DEFAULT_PROFILE); profile["body"][field] = value
    assert client.post(f"/api/uploads/{upload(client)}/apply", json={"profile": profile}).status_code == 422


@pytest.mark.parametrize("ident", ["*", "not-a-uuid", "00000000-0000-0000-0000-000000000000"])
def test_invalid_or_missing_ids_return_404(client, ident):
    assert client.get(f"/api/outputs/{ident}/document").status_code == 404
    assert client.delete(f"/api/uploads/{ident}").status_code == 404


def test_pdf_cannot_be_sent_directly_to_formatter(client):
    import fitz
    doc = fitz.open(); page = doc.new_page(); page.insert_text((60, 60), "Report text")
    response = client.post("/api/uploads", files={"file": ("report.pdf", doc.tobytes())}); doc.close()
    assert response.status_code == 200
    ident = response.json()["id"]
    for action in ("audit", "apply"):
        assert client.post(f"/api/uploads/{ident}/{action}", json={"profile": DEFAULT_PROFILE}).status_code == 400


def test_reference_learning_saves_template(client):
    response = client.post(f"/api/uploads/{upload(client)}/learn-profile")
    assert response.status_code == 200, response.text
    ident = response.json()["profile"]["id"]
    assert ident in [p["id"] for p in client.get("/api/profiles").json()]
    assert client.delete("/api/profiles/generic").status_code == 400
    assert client.delete(f"/api/profiles/{ident}").status_code == 200


def test_alignment_tables_captions_and_emphasis(tmp_path):
    doc = Document(); heading = doc.add_heading("Chapter", 1); heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph(); p.add_run("Important ").bold = True; p.add_run("italic").italic = True
    p.paragraph_format.left_indent = Inches(.7); p.paragraph_format.first_line_indent = Inches(.3)
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    cell = table.cell(1, 0); cell.text = "Table text"; cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    nested = cell.add_table(rows=1, cols=1); nested.cell(0, 0).text = "Nested text"
    caption = doc.add_paragraph("Figure 1. Caption"); caption.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    source = tmp_path / "source.docx"; out = tmp_path / "result.docx"; doc.save(source)
    result = apply_profile(source, out, DEFAULT_PROFILE); clean = Document(out)
    assert clean.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert clean.paragraphs[1].paragraph_format.left_indent == 0
    assert clean.paragraphs[1].paragraph_format.first_line_indent == 0
    assert clean.paragraphs[1].runs[0].bold and clean.paragraphs[1].runs[1].italic
    assert clean.paragraphs[2].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert clean.tables[0].cell(1, 0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert clean.tables[0].cell(0, 0)._tc is clean.tables[0].cell(0, 1)._tc
    assert result["changed"]["tables"] == 2
    assert result["before"] == result["after"]


def test_lists_fields_links_equations_and_breaks_survive(tmp_path):
    doc = Document(); p = doc.add_paragraph("First list item", style="List Number")
    p.paragraph_format.left_indent = Inches(.4); p.paragraph_format.first_line_indent = Inches(-.2)
    num = OxmlElement("w:numPr"); numid = OxmlElement("w:numId"); numid.set(qn("w:val"), "5"); num.append(numid); p._p.get_or_add_pPr().append(num)
    p = doc.add_paragraph("Read ")
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), doc.part.relate_to("https://example.com", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "the source"; r.append(t); hyperlink.append(r); p._p.append(hyperlink)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "REF bookmark"); p._p.append(field)
    math = OxmlElement("m:oMath"); mr = OxmlElement("m:r"); mt = OxmlElement("m:t"); mt.text = "x²"; mr.append(mt); math.append(mr); p._p.append(math)
    doc.add_page_break(); doc.add_paragraph("Second page")
    source = tmp_path / "in.docx"; out = tmp_path / "out.docx"; doc.save(source)
    apply_profile(source, out, DEFAULT_PROFILE); clean = Document(out)
    assert clean.paragraphs[0].paragraph_format.left_indent == Inches(.4)
    assert clean.paragraphs[0].paragraph_format.first_line_indent == Inches(-.2)
    assert signature(Document(source), source) == signature(clean, out)
    assert clean.paragraphs[1]._p.xpath('.//w:hyperlink/w:r/w:rPr/w:rFonts')[0].get(qn("w:ascii")) == "Times New Roman"


def test_fit_images_respects_columns_and_sections(tmp_path):
    image = tmp_path / "pixel.png"
    image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL8xQAAAABJRU5ErkJggg=='))
    doc = Document(); doc.add_picture(str(image), width=Inches(9)); doc.add_paragraph("Illustration for a lesson.")
    section = doc.add_section(WD_SECTION.NEW_PAGE); section.page_width = Inches(11); section.page_height = Inches(8.5)
    doc.add_picture(str(image), width=Inches(9))
    source = tmp_path / "in.docx"; out = tmp_path / "out.docx"; doc.save(source)
    profile = deepcopy(DEFAULT_PROFILE); profile["page"]["columns"] = 2
    apply_profile(source, out, profile); clean = Document(out)
    assert clean.inline_shapes[0].width <= column_width(clean.sections[0])
    assert clean.inline_shapes[1].width <= column_width(clean.sections[1])
    assert clean.inline_shapes[0].width == clean.inline_shapes[0].height
    assert clean.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert clean.sections[1].page_width == Inches(11)


def test_headers_are_not_duplicated_when_linked_or_reformatted(tmp_path):
    doc = Document(); doc.sections[0].header.paragraphs[0].text = "Existing heading"
    doc.add_paragraph("Body"); doc.add_section(WD_SECTION.NEW_PAGE); doc.add_paragraph("Second section")
    source = tmp_path / "in.docx"; first = tmp_path / "first.docx"; second = tmp_path / "second.docx"; doc.save(source)
    profile = deepcopy(DEFAULT_PROFILE); profile["extras"].update(headerText="Added header", footerText="Added footer", pageNumbers=True)
    apply_profile(source, first, profile); apply_profile(first, second, profile)
    clean = Document(second)
    assert [p.text for p in clean.sections[0].header.paragraphs].count("Added header") == 1
    assert clean.sections[0].header.paragraphs[0].text == "Existing heading"
    assert len(clean.sections[0].footer._element.xpath('.//w:fldSimple')) == 1
    assert clean.sections[1].header.is_linked_to_previous


def test_existing_page_field_is_not_duplicated(tmp_path):
    doc = Document(); doc.add_paragraph("Body")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), " PAGE "); doc.sections[0].footer.paragraphs[0]._p.append(field)
    source = tmp_path / "in.docx"; out = tmp_path / "out.docx"; doc.save(source)
    profile = deepcopy(DEFAULT_PROFILE); profile["extras"].update(pageNumbers=True, footerText="Department")
    apply_profile(source, out, profile)
    assert len(Document(out).sections[0].footer._element.xpath('.//w:fldSimple')) == 1


def test_disabled_table_cleanup_preserves_table_format(tmp_path):
    doc = Document(); p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]; p.add_run("Cell").font.name = "Arial"; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    source = tmp_path / "in.docx"; out = tmp_path / "out.docx"; doc.save(source)
    profile = deepcopy(DEFAULT_PROFILE); profile["extras"]["normalizeTables"] = False
    apply_profile(source, out, profile)
    p = Document(out).tables[0].cell(0, 0).paragraphs[0]
    assert p.alignment == WD_ALIGN_PARAGRAPH.RIGHT and p.runs[0].font.name == "Arial"


def test_learning_does_not_treat_exact_line_height_as_multiplier(tmp_path):
    doc = Document(); p = doc.add_paragraph("Example body text"); p.paragraph_format.line_spacing = Pt(14)
    p.runs[0].font.name = "Arial"; p.runs[0].font.size = Pt(11)
    source = tmp_path / "in.docx"; doc.save(source)
    profile = style_profile(source, "Reference")
    assert profile["body"]["font"] == "Arial" and 1 <= profile["body"]["lineSpacing"] <= 3


def test_content_signature_detects_equal_count_text_change(tmp_path):
    a = tmp_path / "a.docx"; b = tmp_path / "b.docx"
    doc = Document(); doc.add_paragraph("Keep this text"); doc.save(a)
    doc.paragraphs[0].text = "Lost some text"; doc.save(b)
    assert metrics(a) == metrics(b)
    assert signature(Document(a), a) != signature(Document(b), b)


def test_formatter_refuses_to_overwrite_original(tmp_path):
    source = tmp_path / 'report.docx'; source.write_bytes(document_bytes()); before = source.read_bytes()
    with pytest.raises(ValueError, match='separate output'):
        apply_profile(source, source, DEFAULT_PROFILE)
    assert source.read_bytes() == before


def test_table_header_text_keeps_contrast(tmp_path):
    from docx.shared import RGBColor
    doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    run = cell.paragraphs[0].add_run('White text on a dark header'); run.font.color.rgb = RGBColor(255, 255, 255)
    shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), '17365D'); cell._tc.get_or_add_tcPr().append(shading)
    source = tmp_path / 'in.docx'; out = tmp_path / 'out.docx'; doc.save(source)
    apply_profile(source, out, DEFAULT_PROFILE)
    assert Document(out).tables[0].cell(0, 0).paragraphs[0].runs[0].font.color.rgb == RGBColor(255, 255, 255)
