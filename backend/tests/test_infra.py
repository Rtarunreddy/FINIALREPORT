from datetime import timedelta
from io import BytesIO

import pytest
from alembic import command
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import inspect

from app.config import Settings, normalize_database_url
from app.db import alembic_config, make_engine
from app.db_models import Base, Document as Doc, User, utcnow
from app.routers import auth as auth_router
from app.services import Cleaner
from app.storage import LocalStorage, S3Storage, check_key
from conftest import build, run_apply, sign_up


def docx_bytes():
    d = Document(); d.add_paragraph("Report."); s = BytesIO(); d.save(s); return s.getvalue()


# ---- migrations ----
def test_migrations_match_models_and_roll_back(tmp_path):
    url = f"sqlite:///{tmp_path / 'm.db'}"; cfg = alembic_config(url)
    command.upgrade(cfg, "head")
    command.check(cfg)  # raises if a model change lacks a migration
    engine = make_engine(url)
    assert set(Base.metadata.tables) <= set(inspect(engine).get_table_names())
    engine.dispose()
    command.downgrade(cfg, "base")
    engine = make_engine(url); assert inspect(engine).get_table_names() == ["alembic_version"]; engine.dispose()


def test_database_url_normalisation():
    assert normalize_database_url("postgres://u:p@h/db") == "postgresql+psycopg://u:p@h/db"
    assert normalize_database_url("postgresql://u:p@h/db") == "postgresql+psycopg://u:p@h/db"
    assert normalize_database_url("sqlite:///x.db") == "sqlite:///x.db"


# ---- storage ----
@pytest.mark.parametrize("key", ["../x", "/abs", "a/../../b", "a\\b", ""])
def test_unsafe_keys_rejected(key):
    with pytest.raises(ValueError): check_key(key)


def test_local_storage_roundtrip(tmp_path):
    store, src = LocalStorage(tmp_path / "s"), tmp_path / "in.bin"; src.write_bytes(b"abc")
    store.put_file("u1/d1", src)
    assert store.read_bytes("u1/d1") == b"abc" and store.exists("u1/d1")
    with store.local_copy("u1/d1") as p: assert p.read_bytes() == b"abc"
    store.delete("u1/d1"); store.delete("u1/d1")
    assert not store.exists("u1/d1")
    with pytest.raises(KeyError): store.read_bytes("u1/d1")
    with pytest.raises(ValueError): store.put_file("../escape", src)


class FakeS3:
    class exceptions:
        class NoSuchKey(Exception): pass
    def __init__(self): self.objects = {}
    def upload_file(self, path, bucket, key): self.objects[(bucket, key)] = open(path, "rb").read()
    def get_object(self, Bucket, Key):
        if (Bucket, Key) not in self.objects: raise self.exceptions.NoSuchKey()
        return {"Body": BytesIO(self.objects[(Bucket, Key)])}
    def delete_object(self, Bucket, Key): self.objects.pop((Bucket, Key), None)
    def head_object(self, Bucket, Key):
        if (Bucket, Key) not in self.objects: raise KeyError(Key)


def test_s3_storage_contract(tmp_path):
    fake = FakeS3(); store = S3Storage(Settings(data_dir=tmp_path, local_mode=True, storage_backend="s3", s3_bucket="b"), client=fake)
    src = tmp_path / "in.bin"; src.write_bytes(b"xyz"); store.put_file("u/d", src)
    assert store.exists("u/d") and store.read_bytes("u/d") == b"xyz"
    with store.local_copy("u/d") as p: assert p.read_bytes() == b"xyz"
    store.delete("u/d")
    with pytest.raises(KeyError): store.read_bytes("u/d")


def test_app_runs_on_s3_storage(tmp_path):
    fake = FakeS3(); settings = Settings(data_dir=tmp_path / "d", session_secret="x" * 40, storage_backend="s3", s3_bucket="b", workers=0)
    from app.main import create_app
    with TestClient(create_app(settings, storage=S3Storage(settings, client=fake))) as base:
        c = sign_up(base.app)
        doc = c.post("/api/uploads", files={"file": ("r.docx", docx_bytes())}).json()["id"]
        from app.formatting import DEFAULT_PROFILE
        out = run_apply(c, doc, DEFAULT_PROFILE)["id"]
        assert len(fake.objects) == 2 and c.get(f"/api/outputs/{out}/document").status_code == 200
        c.delete(f"/api/uploads/{doc}"); assert fake.objects == {}


# ---- retention ----
def test_expired_documents_are_removed_and_unreachable(cloud):
    c = sign_up(cloud.app)
    doc = c.post("/api/uploads", files={"file": ("r.docx", docx_bytes())}).json()["id"]
    from conftest import stored_files
    assert len(stored_files(cloud)) == 1
    with cloud.app.state.session_factory() as db:
        row = db.get(Doc, doc); assert row.expires_at - row.created_at == timedelta(hours=24)
        row.expires_at = utcnow() - timedelta(minutes=1); db.commit()
    from app.formatting import DEFAULT_PROFILE
    assert c.post(f"/api/uploads/{doc}/audit", json={"profile": DEFAULT_PROFILE}).status_code == 404   # blocked even before the sweep
    assert Cleaner().run(cloud.app.state.session_factory, cloud.app.state.storage, force=True) == 1
    assert stored_files(cloud) == []
    with cloud.app.state.session_factory() as db: assert db.get(Doc, doc).deleted_at is not None


def test_paid_plans_keep_history_longer(cloud):
    c = sign_up(cloud.app)
    with cloud.app.state.session_factory() as db:
        user = db.query(User).one()
        from app.db_models import Subscription
        db.add(Subscription(owner_type="user", owner_id=user.id, stripe_sub_id="sub_paid", plan="pro", status="active", current_period_end=utcnow()+timedelta(days=30))); db.commit()
    assert c.get("/api/auth/status").json()["retentionHours"] == 24 * 90
    doc = c.post("/api/uploads", files={"file": ("r.docx", docx_bytes())}).json()["id"]
    with cloud.app.state.session_factory() as db:
        row = db.get(Doc, doc); assert row.expires_at - row.created_at == timedelta(hours=24 * 90)


def test_storage_keys_never_contain_user_filenames(cloud):
    c = sign_up(cloud.app)
    c.post("/api/uploads", files={"file": ("../../evil name.docx", docx_bytes())})
    from conftest import stored_files
    (path,) = stored_files(cloud)
    assert "evil" not in str(path.relative_to(cloud.app.state.settings.data_dir / "files"))


# ---- Google sign-in ----
@pytest.fixture
def google(tmp_path):
    with TestClient(build(tmp_path, google_client_id="cid", google_client_secret="sec", public_url="https://app.example")) as c: yield c


def callback(c, state="s1", cookie_state="s1"):
    c.cookies.set("report_ready_oauth_state", cookie_state)
    return c.get("/api/auth/google/callback", params={"code": "abc", "state": state}, follow_redirects=False)


def test_google_login_redirects_with_state(google):
    r = google.get("/api/auth/google/login", follow_redirects=False)
    assert r.status_code in (302, 307) and "accounts.google.com" in r.headers["location"] and "state=" in r.headers["location"]
    assert google.get("/api/auth/status").json()["googleEnabled"] is True


def test_google_disabled_by_default(cloud):
    assert cloud.get("/api/auth/google/login").status_code == 404


def test_google_callback_creates_and_links_accounts(google, monkeypatch):
    monkeypatch.setattr(auth_router, "exchange_google_code", lambda s, code: {"sub": "g-1", "email": "Pat@Example.com", "email_verified": True})
    assert callback(google).status_code == 303 and google.get("/api/auth/me").json()["email"] == "pat@example.com"
    with google.app.state.session_factory() as db: assert db.query(User).one().google_sub == "g-1"
    # an existing password account is linked, not duplicated, when Google vouches for the email
    other = TestClient(google.app); other.post("/api/auth/register", json={"email": "lee@example.com", "password": "correct horse battery"})
    monkeypatch.setattr(auth_router, "exchange_google_code", lambda s, code: {"sub": "g-2", "email": "lee@example.com", "email_verified": True})
    assert callback(TestClient(google.app)).status_code == 303
    with google.app.state.session_factory() as db: assert db.query(User).count() == 2


@pytest.mark.parametrize("info", [{"sub": "g", "email": "a@example.com", "email_verified": False}, {"sub": "g", "email": "a@example.com"}, {"email": "a@example.com", "email_verified": True}])
def test_google_requires_verified_email(google, monkeypatch, info):
    monkeypatch.setattr(auth_router, "exchange_google_code", lambda s, code: info)
    assert callback(google).status_code == 400


def test_google_state_mismatch_rejected(google, monkeypatch):
    monkeypatch.setattr(auth_router, "exchange_google_code", lambda s, code: {"sub": "g", "email": "a@example.com", "email_verified": True})
    assert callback(google, state="evil", cookie_state="s1").status_code == 400
    assert TestClient(google.app).get("/api/auth/google/callback", params={"code": "abc", "state": "s1"}).status_code == 400
