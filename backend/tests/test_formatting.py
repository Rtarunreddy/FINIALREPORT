from pathlib import Path
import base64
from docx import Document
from docx.shared import Inches, Pt
from app.formatting import APA_PROFILE, CHICAGO_PROFILE, DEFAULT_PROFILE, HARVARD_PROFILE, IEEE_PROFILE, MLA_PROFILE, apply_profile, metrics
from app.routers.templates import BUILTINS
from app.models import Profile

def test_formats_safe_body_and_preserves_content(tmp_path):
    source=tmp_path/'input.docx'; out=tmp_path/'output.docx'; doc=Document()
    doc.add_heading('Chapter one', 1); p=doc.add_paragraph('This is the report body text.')
    p.runs[0].font.name='Arial'; p.runs[0].font.size=Pt(9)
    listed=doc.add_paragraph('This list content should be formatted.', style='List Paragraph')
    listed.runs[0].font.name='Arial'; listed.runs[0].font.size=Pt(9)
    doc.add_table(rows=1, cols=1).cell(0,0).text='A table stays here'
    sec=doc.sections[0]; sec.orientation=1; sec.page_width,sec.page_height=Inches(11.7), Inches(8.3)
    doc.save(source); before=metrics(source)
    report=apply_profile(source,out,DEFAULT_PROFILE)
    result=Document(out); body=result.paragraphs[1]
    assert body.runs[0].font.name == 'Times New Roman'
    assert round(body.runs[0].font.size.pt) == 12
    assert result.paragraphs[2].runs[0].font.name == 'Times New Roman'
    assert metrics(out)==before and report['changed']['body']==1 and report['changed']['lists']==1
    assert result.tables[0].cell(0,0).text=='A table stays here'

def test_resizes_inline_images_without_removing_them(tmp_path):
    source=tmp_path/'input.docx'; out=tmp_path/'output.docx'; image=tmp_path/'pixel.png'
    image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL8xQAAAABJRU5ErkJggg=='))
    doc=Document(); doc.add_paragraph('Body text'); doc.add_picture(str(image), width=Inches(8)); doc.save(source)
    profile={**DEFAULT_PROFILE, 'extras': {'imageMaxWidth': 3}}
    report=apply_profile(source,out,profile); result=Document(out)
    assert report['changed']['images']==1 and result.inline_shapes[0].width <= Inches(3)
    assert metrics(source)['images']==metrics(out)['images']==1

def test_academic_presets_apply_standard_indents_and_page_numbers(tmp_path):
    source=tmp_path/'input.docx'; out=tmp_path/'output.docx'
    doc=Document(); doc.add_heading('Chapter one', 1); doc.add_paragraph('A paragraph with a standard academic indent.')
    doc.save(source)
    report=apply_profile(source,out,APA_PROFILE); result=Document(out)
    assert result.sections[0].page_width == Inches(8.5)
    assert result.paragraphs[1].paragraph_format.first_line_indent == Inches(.5)
    assert result.paragraphs[1].paragraph_format.line_spacing == 2
    assert result.sections[0].header._element.xpath('.//w:fldSimple[@w:instr="PAGE"]')
    assert report['changed']['pageNumbers'] == 1
    assert all(p['body']['firstLineIndent'] == .5 for p in (APA_PROFILE, MLA_PROFILE, CHICAGO_PROFILE, HARVARD_PROFILE))
    assert IEEE_PROFILE['page']['columns'] == 2

def test_page_border_box_and_double_styles_are_written_to_each_section(tmp_path):
    source=tmp_path/'input.docx'; out=tmp_path/'output.docx'
    Document().save(source)
    for style, expected in (('box', 'single'), ('double', 'double')):
        profile={**DEFAULT_PROFILE, 'page': {**DEFAULT_PROFILE['page'], 'border': style}}
        apply_profile(source,out,profile)
        result=Document(out)
        borders=result.sections[0]._sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgBorders')
        assert borders is not None
        assert [edge.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') for edge in borders] == [expected]*4

def test_all_builtin_format_profiles_validate_and_are_selectable():
    names={item['id']: Profile.model_validate(item).name for item in BUILTINS}
    assert names == {
        'generic': 'Generic university report',
        'technical': 'Formal technical report',
        'apa7': 'APA 7 (student-paper layout)',
        'mla9': 'MLA 9 (paper layout)',
        'chicago': 'Chicago / Turabian (student-paper layout)',
        'harvard': 'Harvard author-date (institution-dependent)',
        'ieee': 'IEEE (two-column paper layout)',
        'business': 'Business report',
    }
