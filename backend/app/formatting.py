"""Format DOCX in place without rebuilding its content or drawing relationships."""
import hashlib
import re
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from .models import Profile

DEFAULT_PROFILE = Profile(name="Generic university report").model_dump(exclude={"id"})
TECHNICAL_PROFILE = deepcopy(DEFAULT_PROFILE)
TECHNICAL_PROFILE["name"] = "Formal technical report"
TECHNICAL_PROFILE["page"]["columns"] = 1
IEEE_PROFILE = deepcopy(DEFAULT_PROFILE)
IEEE_PROFILE["name"] = "Two-column paper starter"
IEEE_PROFILE["page"].update(top=.75, bottom=.75, left=.75, right=.75, columns=2)
IEEE_PROFILE["body"].update(size=10, lineSpacing=1, after=0)
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def styles(p):
    style = p.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def heading_level(p):
    for style in styles(p):
        name = style.name.lower().replace(" ", "")
        if name in {"heading1", "heading2", "heading3"}:
            return int(name[-1])
    nodes = p._p.xpath("./w:pPr/w:outlineLvl")
    return int(nodes[0].get(qn("w:val"))) + 1 if nodes and nodes[0].get(qn("w:val")) in {"0", "1", "2"} else None


def is_list(p):
    return bool(p._p.xpath("./w:pPr/w:numPr")) or any(
        s.name.lower().startswith("list") or s.element.xpath("./w:pPr/w:numPr") for s in styles(p))


def role(p, in_table=False):
    # Only named, known roles are changed. Unknown custom styles remain review items.
    names = {s.name.lower() for s in styles(p)}
    own = p.style.name.lower() if p.style is not None else "normal"
    if heading_level(p): return "headings"
    if own == "title": return "titles"
    if own == "subtitle": return "subtitles"
    previous = p._p.getprevious()
    adjacent_illustration = (previous is not None and previous.tag == qn("w:p") and previous.xpath(".//wp:inline")
                            and len(p.text) < 300 and re.match(r"^\s*Illustration\b", p.text, re.I))
    if "caption" in names or adjacent_illustration or re.match(r"^\s*(?:figure|fig\.|table)\s+\d+[\w.\-]*\s*[:.\-–—]", p.text, re.I): return "captions"
    if is_list(p): return "lists"
    if in_table: return "tableText"
    if own in {"normal", "normal (web)", "body text", "bodytext"}: return "body"
    return "custom"


def all_paragraphs(doc):
    for element in doc.element.body.iter(qn("w:p")):
        # Text inside a drawing is left alone: its available geometry is unrelated.
        if any(a.tag in {qn("w:txbxContent"), qn("w:drawing")} for a in element.iterancestors()): continue
        yield Paragraph(element, doc)


def body_paragraphs(doc):
    return [p for p in doc.paragraphs if p.text.strip() and role(p) in {"body", "lists"}]


def runs(p):
    for element in p._p.iter(qn("w:r")):
        if next(element.iterancestors(qn("w:p")), None) is p._p:
            yield Run(element, p)


def metrics(path: Path):
    if path.suffix.lower() == ".pdf":
        import fitz
        with fitz.open(path) as pdf:
            return {"pages": len(pdf), "words": sum(len(p.get_text().split()) for p in pdf), "images": sum(len(p.get_images()) for p in pdf)}
    doc = Document(path)
    paragraphs = list(doc.element.body.iter(qn("w:p")))
    with zipfile.ZipFile(path) as z:
        images = sum(n.startswith("word/media/") and not n.endswith("/") for n in z.namelist())
    return {"paragraphs": len(doc.paragraphs),
            "images": images, "tables": len(list(doc.element.body.iter(qn("w:tbl")))),
            "allParagraphs": len(paragraphs),
            "words": sum(len("".join(t.text or "" for t in p.findall(".//" + qn("w:t")) if next(t.iterancestors(qn("w:p")), None) is p).split()) for p in paragraphs),
            "imageOccurrences": len(doc.element.body.xpath(".//wp:inline | .//wp:anchor"))}


def signature(doc, path):
    """Verify actual content, rather than accepting equal word counts as proof."""
    body = doc.element.body
    tags = ["w:t", "w:instrText", "w:delText", "w:tab", "w:br", "w:fldChar", "w:fldSimple", "w:footnoteReference", "w:endnoteReference", "m:t", "a:blip"]
    allowed = {qn(t) for t in tags}
    content = [(e.tag, e.text, sorted(e.attrib.items())) for e in body.iter() if e.tag in allowed]
    structure = [e.tag for e in body.iter() if e.tag in {qn(x) for x in ["w:p", "w:tbl", "w:tr", "w:tc", "w:hyperlink", "m:oMath", "w:sectPr"]}]
    links = ([sorted(e.attrib.items()) for e in body.iter(qn("w:hyperlink"))],
             sorted((r.rId, r.target_ref) for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")))
    numbering = [etree.tostring(e, method="c14n") for e in body.xpath(".//w:numPr")]
    equations = [etree.tostring(e, method="c14n") for e in body.xpath(".//m:oMath")]
    with zipfile.ZipFile(path) as z:
        assets = {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith(("word/media/", "word/embeddings/"))}
    return content, structure, links, numbering, equations, assets


def format_text(p, settings, reset_indent=True, bold=None):
    fmt = p.paragraph_format
    fmt.alignment = ALIGN[settings.get("alignment", "left")]
    fmt.line_spacing = settings.get("lineSpacing", 1.15)
    fmt.space_before = Pt(settings.get("before", 0))
    fmt.space_after = Pt(settings.get("after", 6))
    fmt.widow_control = True
    if reset_indent:
        fmt.left_indent = Pt(0); fmt.right_indent = Pt(0); fmt.first_line_indent = Pt(0)
        ind = p._p.get_or_add_pPr().find(qn("w:ind"))
        if ind is not None:
            for key in ("leftChars", "rightChars", "firstLineChars", "hangingChars", "start", "end"):
                ind.attrib.pop(qn("w:" + key), None)
    for r in runs(p):
        # Drawing-only runs and native equation children are not restyled.
        if not r.text: continue
        r.font.name = settings["font"]; r.font.size = Pt(settings["size"])
        fonts = r._r.get_or_add_rPr().rFonts
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + key), None)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"): fonts.set(qn("w:" + key), settings["font"])
        if not settings.get("preserveColor") and not any(a.tag == qn("w:hyperlink") for a in r._r.iterancestors()):
            r.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None: r.bold = bold


def section_for(element, sections):
    top = element
    while top.getparent() is not None and top.getparent().tag != qn("w:body"): top = top.getparent()
    current = top
    while current is not None:
        candidates = [current] if current.tag == qn("w:sectPr") else current.xpath("./w:pPr/w:sectPr")
        if candidates:
            return next((s for s in sections if s._sectPr is candidates[0]), sections[-1])
        current = current.getnext()
    return sections[-1]


def column_width(section):
    width = section.page_width - section.left_margin - section.right_margin
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None: return width
    explicit = cols.findall(qn("w:col"))
    if explicit:
        return min(int(c.get(qn("w:w"), "1")) * 635 for c in explicit)
    count = int(cols.get(qn("w:num"), "1"))
    gap = int(cols.get(qn("w:space"), "720")) * 635
    return (width - gap * (count - 1)) / count


def fit_tables(doc, changed):
    for table in doc.element.body.iter(qn("w:tbl")):
        pr = table.find(qn("w:tblPr"))
        if pr is None: pr = OxmlElement("w:tblPr"); table.insert(0, pr)
        for tag in ("w:jc", "w:tblInd"):
            for old in pr.findall(qn(tag)): pr.remove(old)
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); pr.append(jc)
        # Preserve merged cells and relative column widths. Only shrink overflow.
        width = column_width(section_for(table, doc.sections)) / 635
        parent_cell = next(table.iterancestors(qn("w:tc")), None)
        if parent_cell is not None:
            cw = parent_cell.find("./" + qn("w:tcPr") + "/" + qn("w:tcW"))
            if cw is not None and cw.get(qn("w:type")) == "dxa": width = min(width, int(cw.get(qn("w:w"))) - 216)
        grid = table.findall("./" + qn("w:tblGrid") + "/" + qn("w:gridCol"))
        total = sum(int(c.get(qn("w:w"), "0")) for c in grid)
        ratio = min(1, max(1, width) / total) if total else 1
        if ratio < 1:
            for c in grid: c.set(qn("w:w"), str(max(1, round(int(c.get(qn("w:w"))) * ratio))))
            for cell in table.xpath("./w:tr/w:tc/w:tcPr/w:tcW"):
                if cell.get(qn("w:type")) == "dxa": cell.set(qn("w:w"), str(max(1, round(int(cell.get(qn("w:w"))) * ratio))))
            tw = pr.find(qn("w:tblW"))
            if tw is not None: tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(round(width)))
        for row_height in table.xpath("./w:tr/w:trPr/w:trHeight"):
            if row_height.get(qn("w:hRule")) == "exact": row_height.set(qn("w:hRule"), "atLeast")
        changed["tables"] += 1


def has_page_number(container):
    return any(re.search(r"\bPAGE\b", text, re.I) for text in
               [e.text or "" for e in container._element.iter(qn("w:instrText"))] +
               [e.get(qn("w:instr"), "") for e in container._element.iter(qn("w:fldSimple"))])


def managed_paragraph(container, name, text, page=False):
    matches = container._element.xpath(f'.//w:p[w:bookmarkStart[@w:name="{name}"]]')
    p = Paragraph(matches[0], container) if matches else container.add_paragraph()
    # This paragraph belongs to this app. Existing user paragraphs are never replaced.
    p.clear()
    marker = OxmlElement("w:bookmarkStart"); marker.set(qn("w:id"), "2147483600"); marker.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), "2147483600")
    p._p.append(marker); p.add_run(text)
    if page:
        if text: p.add_run(" · ")
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); p._p.append(field)
    p._p.append(end)
    return p


def format_headers(doc, profile, changed):
    extras = profile["extras"]; seen = set()
    settings = {**profile["body"], "size": 9, "lineSpacing": 1, "before": 0, "after": 0}
    for section in doc.sections:
        for kind in ("header", "footer"):
            variants = [getattr(section, kind)]
            for prefix in ("first_page_", "even_page_"):
                variant = getattr(section, prefix + kind)
                if variant._has_definition: variants.append(variant)
            for container in variants:
                text = extras[kind + "Text"].strip()
                want_page = extras["pageNumbers"] and extras["pageNumberPosition"] == kind
                if not container._has_definition and container.is_linked_to_previous and not text and not want_page: continue
                element = container._element
                if element in seen: continue
                seen.add(element)
                marker_name = "ReportReady" + kind.title()
                owned = element.xpath(f'.//w:p[w:bookmarkStart[@w:name="{marker_name}"]]')
                # Updating our own field should not be mistaken for a user PAGE field.
                existing_page = any(has_page_number(ParagraphContainer(p)) for p in element.findall(qn("w:p")) if p not in owned)
                if text or want_page:
                    managed_paragraph(container, marker_name, text, want_page and not existing_page)
                    if want_page and not existing_page: changed["pageNumbers"] += 1
                if extras["normalizeHeadersFooters"]:
                    for p in container.paragraphs:
                        format_text(p, {**settings, "alignment": "left" if kind == "header" else "center"})
                if text or want_page or extras["normalizeHeadersFooters"]: changed[kind + "s"] += 1


class ParagraphContainer:
    def __init__(self, element): self._element = element


def apply_profile(source: Path, dest: Path, profile: dict):
    if source.resolve() == dest.resolve(): raise ValueError("Choose a separate output file so the original stays unchanged.")
    profile = Profile.model_validate(profile).model_dump()
    doc = Document(source); before = metrics(source); original = signature(doc, source)
    changed = {k: 0 for k in ("body", "lists", "headings", "titles", "subtitles", "captions", "tableText", "tables", "sections", "headers", "footers", "pageNumbers", "images", "imageParagraphs")}
    page = profile["page"]; extras = profile["extras"]; body = profile["body"]
    for section in doc.sections:
        if page["size"] != "Existing report":
            w, h = (8.5, 11) if page["size"] == "Letter" else (8.2677, 11.6929)
            if section.page_width > section.page_height: w, h = h, w
            section.page_width = Inches(w); section.page_height = Inches(h)
        for side in ("top", "bottom", "left", "right"): setattr(section, side + "_margin", Inches(page[side]))
        if section.page_width - section.left_margin - section.right_margin < Inches(1) or section.page_height - section.top_margin - section.bottom_margin < Inches(1):
            raise ValueError("The margins leave too little writing space. Choose smaller margins.")
        if page["columns"]:
            cols = section._sectPr.find(qn("w:cols"))
            if cols is None: cols = OxmlElement("w:cols"); section._sectPr.append(cols)
            for child in list(cols): cols.remove(child)
            cols.set(qn("w:num"), str(page["columns"])); cols.set(qn("w:equalWidth"), "1"); cols.set(qn("w:space"), "360")
        changed["sections"] += 1
    if extras["normalizeTables"]: fit_tables(doc, changed)
    for p in all_paragraphs(doc):
        in_table = next(p._p.iterancestors(qn("w:tc")), None) is not None
        kind = role(p, in_table)
        if in_table and not extras["normalizeTables"]: continue
        pictures = p._p.xpath(".//wp:inline")
        if pictures and not p.text.strip() and extras["centerImages"]:
            format_text(p, {**body, "alignment": "center", "before": 6, "after": 6, "lineSpacing": 1})
            p.paragraph_format.keep_with_next = bool(p._p.getnext() is not None and p._p.getnext().tag == qn("w:p") and role(Paragraph(p._p.getnext(), doc)) == "captions")
            changed["imageParagraphs"] += 1
        if not p.text.strip(): continue
        settings = dict(body); reset = extras["resetBodyIndents"]; bold = None
        if kind == "custom": continue
        if kind == "headings":
            settings = profile["headings"][f"h{heading_level(p)}"]; bold = settings["bold"]; reset = not is_list(p)
            p.paragraph_format.keep_with_next = True
            if heading_level(p) == 1 and extras["startChaptersOnNewPage"]: p.paragraph_format.page_break_before = True
        elif kind in {"titles", "subtitles"}:
            settings.update(size=24 if kind == "titles" else 14, alignment="left", lineSpacing=1.15, before=0, after=12)
            bold = kind == "titles"; reset = True; p.paragraph_format.keep_with_next = True
        elif kind == "captions":
            if not extras["normalizeCaptions"]: continue
            settings.update(size=max(9, body["size"] - 2), alignment="center", lineSpacing=1, before=4, after=8); reset = True
        elif kind == "lists": reset = False; settings["alignment"] = "left"
        elif kind == "tableText": settings.update(size=min(11, body["size"]), alignment="left", lineSpacing=1.1, before=0, after=4); reset = True
        if in_table: settings = {**settings, "preserveColor": True}
        format_text(p, settings, reset, bold); changed[kind] += 1
    for shape in doc.inline_shapes:
        section = section_for(shape._inline, doc.sections)
        width = column_width(section)
        cell = next(shape._inline.iterancestors(qn("w:tc")), None)
        if cell is not None:
            cw = cell.find("./" + qn("w:tcPr") + "/" + qn("w:tcW"))
            if cw is not None and cw.get(qn("w:type")) == "dxa": width = min(width, max(1, int(cw.get(qn("w:w"))) - 216) * 635)
        limits = [1.0]
        if extras["fitImages"]:
            limits.extend([width / max(1, shape.width), (section.page_height - section.top_margin - section.bottom_margin - Inches(.5)) / max(1, shape.height)])
        if extras["imageMaxWidth"]: limits.append(Inches(extras["imageMaxWidth"]) / max(1, shape.width))
        ratio = min(limits)
        if ratio < 1:
            shape.width = int(shape.width * ratio); shape.height = int(shape.height * ratio); changed["images"] += 1
    format_headers(doc, profile, changed)
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(dest)
        after = metrics(dest)
        if original != signature(Document(dest), dest) or before != after:
            raise ValueError("The content check found an unexpected change. No output was kept.")
    except Exception:
        dest.unlink(missing_ok=True); raise
    return {"before": before, "after": after, "changed": changed,
            "formatting_preview": {"body": f"{body['font']} {body['size']:g} pt · {body['alignment']} · {body['lineSpacing']:g} line spacing",
                "header": extras["headerText"] or "Existing header text retained",
                "footer": extras["footerText"] or "Existing footer text retained",
                "images": "Inline pictures fit within the page or column, keeping their proportions" if extras["fitImages"] else "Image fitting disabled"},
            "preserved": ["Original upload unchanged", "Body text, fields, lists, links, equations and embedded asset bytes verified", "Existing header and footer content retained"],
            "manual": "Open the DOCX in Word to review page breaks and complex layouts. Page count may change when formatting changes."}


def audit(path: Path, profile: dict):
    profile = Profile.model_validate(profile).model_dump(); doc = Document(path)
    counts = Counter(role(p, next(p._p.iterancestors(qn("w:tc")), None) is not None) for p in all_paragraphs(doc) if p.text.strip())
    issues = []; review = []
    if counts["custom"]: review.append({"title": f"{counts['custom']} paragraphs use custom styles", "why": "Their role is uncertain.", "action": "They will retain their current formatting. Use Heading 1–3 for headings you want normalized."})
    if doc.element.body.xpath(".//wp:anchor | .//w:txbxContent | .//w:ins | .//w:del"):
        review.append({"title": "Complex content needs a visual check", "why": "Floating objects, text boxes or tracked changes are present.", "action": "Review their placement in Word after formatting."})
    for key, label in (("body", "Body text"), ("headings", "Headings"), ("lists", "List text"), ("titles", "Titles")):
        if counts[key]: issues.append({"title": f"{label}: {counts[key]} paragraphs", "why": "Apply the selected template consistently.", "action": "Normalize font, spacing and alignment; retain list numbering and emphasis.", "severity": "fixable"})
    for enabled, key, label in (("normalizeTables", "tableText", "Table text"), ("normalizeCaptions", "captions", "Captions")):
        if counts[key] and profile["extras"][enabled]: issues.append({"title": f"{label}: {counts[key]} paragraphs", "why": "These were missed by the earlier formatter.", "action": "Apply consistent font, spacing and alignment.", "severity": "fixable"})
    return {"passed": ["DOCX opened successfully", "A separate formatted copy will be created"], "issues": issues, "review": review, "metrics": metrics(path),
            "warning": "This is a formatting plan, not a compliance certificate. Pagination can change; review the downloaded document in Word."}


def style_profile(path: Path, name: str):
    doc = Document(path); result = deepcopy(DEFAULT_PROFILE); result["name"] = name
    sec = doc.sections[0]
    result["page"].update({side: round(getattr(sec, side + "_margin").inches, 2) for side in ("top", "bottom", "left", "right")})
    def effective(obj, attr, p):
        value = getattr(obj, attr)
        if value is not None: return value
        for s in styles(p):
            value = getattr(s.font if attr in {"name", "size", "bold"} else s.paragraph_format, attr)
            if value is not None: return value
        return None
    def learn(paragraphs, settings):
        if not paragraphs: return
        for attr, key in (("name", "font"), ("size", "size")):
            values = Counter()
            for p in paragraphs:
                for r in runs(p):
                    value = effective(r.font, attr, p)
                    if value is not None and r.text.strip(): values[value.pt if attr == "size" else value] += len(r.text)
            if values: settings[key] = values.most_common(1)[0][0]
        for attr, key in (("alignment", "alignment"), ("line_spacing", "lineSpacing"), ("space_before", "before"), ("space_after", "after")):
            values = Counter(effective(p.paragraph_format, attr, p) for p in paragraphs)
            value = values.most_common(1)[0][0]
            if value is None: continue
            if attr == "alignment": value = next((k for k, v in ALIGN.items() if v == value), "left")
            elif attr == "line_spacing":
                if not isinstance(value, float): continue  # Exact points are not a line multiplier.
            else: value = value.pt
            settings[key] = round(value, 2) if isinstance(value, float) else value
    learn([p for p in doc.paragraphs if role(p) == "body"], result["body"])
    for level in (1, 2, 3): learn([p for p in doc.paragraphs if heading_level(p) == level], result["headings"][f"h{level}"])
    return Profile.model_validate(result).model_dump(exclude={"id"})
